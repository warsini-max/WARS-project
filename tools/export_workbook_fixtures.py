"""Export synthetic test pairs from LibreOffice-recalculated WARS Track B v0.6 builds.

The workbook (WARS_TrackB_v0_6.xlsx) is the specification of the calculation engine.
Earlier work produced test builds of that workbook filled with SYNTHETIC pairs
(tickers such as TST01, TJ01, TK01) and recalculated them in LibreOffice. This script
copies, for those synthetic rows only:
  - every input cell of Pair_Inputs, and
  - the values LibreOffice calculated in Pair_Inputs (derived columns) and in
    Beneish_Pair_Calculation,
into a JSON fixture that the engine tests compare against.

Rows of real issuers are NOT exported: real case data
stays out of the repository, and evaluation cases are not used before the
evaluation rules are locked (plan r0, Phase 7).

Usage:
    python tools/export_workbook_fixtures.py <recalc.xlsx> [<recalc.xlsx> ...] \
        --out tests/fixtures/workbook_v06_synthetic.json
"""
import argparse
import datetime as dt
import hashlib
import json
import re
from pathlib import Path

import openpyxl

SYNTHETIC = re.compile(r"^T[A-Z]{1,3}\d{2}[A-Z]?$")  # TST01, TSH01, TSI07B, TJ01, TK01, TL01, TMA00

PAIR_DERIVED = [
    "obs_key", "row_active", "is_primary", "basis_specified", "gross_profit_t", "gross_profit_t1",
    "ppe_dep_check_t", "ppe_dep_check_t1", "depreciation_expense_t", "depreciation_expense_t1",
    "required_inputs_complete_t", "required_inputs_complete_t1",
]


def jsonable(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()
    return v


def rows(ws, header_row, first_data_row):
    header = list(next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True)))
    out = []
    for r in ws.iter_rows(min_row=first_data_row, values_only=True):
        out.append(dict(zip(header, r)))
    return header, out


def export(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    _, pairs = rows(wb["Pair_Inputs"], 5, 6)
    _, calc = rows(wb["Beneish_Pair_Calculation"], 5, 6)
    cases = []
    for i, p in enumerate(pairs):
        t = p.get("ticker")
        if not t or not SYNTHETIC.match(str(t)):
            continue
        c = calc[i]
        assert c.get("pair_row") == i + 1, (path, i, c.get("pair_row"))
        inputs = {k: jsonable(v) for k, v in p.items() if k and k not in PAIR_DERIVED and v is not None}
        expected_pair = {k: jsonable(p.get(k)) for k in PAIR_DERIVED}
        expected_calc = {k: jsonable(v) for k, v in c.items() if k}
        cases.append({"source_build": Path(path).name, "pair_id": p.get("pair_id"),
                      "inputs": inputs, "expected_pair_inputs": expected_pair,
                      "expected_calculation": expected_calc})
    return cases


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("builds", nargs="+")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    all_cases, sources = [], []
    for b in a.builds:
        cs = export(b)
        all_cases += cs
        sources.append({"file": Path(b).name, "sha256": hashlib.sha256(Path(b).read_bytes()).hexdigest(),
                        "synthetic_rows": len(cs)})
    doc = {"description": "Synthetic Track B pairs and LibreOffice-recalculated values from WARS_TrackB_v0_6 test builds",
           "sources": sources, "cases": all_cases}
    Path(a.out).write_text(json.dumps(doc, indent=1, ensure_ascii=False))
    print(f"{len(all_cases)} synthetic cases from {len(sources)} builds -> {a.out}")


if __name__ == "__main__":
    main()
