"""The Excel download (Gail 28 Oct 2026: replaces the CSV with a tidy table), in ID or EN.

Sheets: Summary · Figures (every value read, both years, pages, source line) · Indices ·
Checklist (with a Matches / Does not match drop-down) · Cross-checks. Amounts are stored as numbers
in the report's unit, so they can be recalculated; blanks stay blank (never 0).
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import report
from .text import render, t, unit_msg

HEAD_FILL = PatternFill("solid", fgColor="006B53")
HEAD_FONT = Font(name="Calibri", bold=True, color="FFFFFF")
FILL = {report.GREEN: None, report.YELLOW: PatternFill("solid", fgColor="FDF4DC"),
        report.RED: PatternFill("solid", fgColor="F9E1E0")}
THIN = Border(bottom=Side(style="thin", color="D9DCD8"))
WRAP = Alignment(wrap_text=True, vertical="top")
NUM = "#,##0;(#,##0)"


def _sheet(wb, title, header, widths):
    ws = wb.create_sheet(title[:31])
    ws.append(header)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
        c = ws.cell(row=1, column=j)
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    return ws


def _finish(ws, num_cols=()):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment, c.border = WRAP, THIN
            if c.column in num_cols and isinstance(c.value, (int, float)):
                c.number_format = NUM
    ws.auto_filter.ref = ws.dimensions


def build(r, lang="id", file_name="", sha="", generated="") -> bytes:
    ex = r.extraction
    fy = ex.fiscal_year
    yt, yt1 = (str(fy), str(fy - 1)) if fy else ("t", "t−1")
    ticker, name = report.issuer_line(r, lang)
    wb = Workbook()
    wb.remove(wb.active)

    # ---------------------------------------------------------------- summary
    ws = _sheet(wb, t("xl.sheet.summary", lang), ["", ""], [34, 90])
    head, detail, _ = report.verdict(r, None, lang, provisional_note=False)
    unit = unit_msg(ex.unit_scale, getattr(ex, "currency", None))
    stype = t("ui.res.stype.Consolidated" if ex.statement_type == "Consolidated" else "ui.res.stype.unknown", lang)
    rows = [(t("rep.f.entity", lang), name), (t("rep.f.ticker", lang), ticker or "—"),
            (t("rep.f.period", lang), t("rep.period", lang, t=yt, t1=yt1)), (t("rep.f.basis", lang), stype),
            (t("rep.f.unit", lang), render(unit, lang)), (t("xl.k.file", lang), file_name), ("SHA-256", sha),
            (t("xl.k.mscore", lang), round(r.m_score, 3) if r.m_score is not None else "—"),
            (t("xl.k.status", lang), report.status_text(r, None, lang)), (t("rep.summary", lang), f"{head}. {detail}"),
            (t("xl.k.generated", lang), generated), ("", t("app.disclaimer", lang))]
    for k, v in rows:
        ws.append([k, v])
    ws.delete_rows(1)
    for row in ws.iter_rows():
        row[0].font = Font(bold=True)
        for c in row:
            c.alignment = WRAP

    # ---------------------------------------------------------------- figures
    by = {(v.field, v.column): v for v in ex.values}
    hdr = [t("inputs.col.item", lang), yt, yt1, t("inputs.col.status", lang, y=yt), t("inputs.col.status", lang, y=yt1),
           t("inputs.col.pages", lang), t("xl.col.source_label", lang), t("xl.col.source_line", lang),
           t("inputs.col.note", lang)]
    ws = _sheet(wb, t("xl.sheet.figures", lang), hdr, [44, 18, 18, 26, 26, 12, 30, 60, 60])
    for x in report.input_rows(r, lang):
        vt, vt1 = by.get((x["field"], "t")), by.get((x["field"], "t1"))
        num = lambda v: v.value if v is not None and isinstance(v.value, (int, float)) else None   # noqa: E731
        ws.append([x["label"].strip(), num(vt), num(vt1), x["status_t"], x["status_t1"], x["pages"],
                   report.source_label(x["source"], lang), x["source"], x["note"]])
        i = ws.max_row
        for col, colour in ((2, x["colour_t"]), (3, x["colour_t1"])):
            if FILL[colour] is not None:
                ws.cell(row=i, column=col).fill = FILL[colour]
    _finish(ws, num_cols=(2, 3))

    # ---------------------------------------------------------------- indices
    hdr = [t("idx.col.index", lang), t("idx.col.name", lang), t("idx.col.value", lang), t("idx.col.coef", lang),
           t("idx.col.contribution", lang), t("idx.col.push", lang), t("idx.col.note", lang)]
    ws = _sheet(wb, t("xl.sheet.indices", lang), hdr, [10, 40, 12, 12, 14, 14, 70])
    for x in report.index_rows(r, lang):
        rnd = lambda v: round(v, 4) if isinstance(v, (int, float)) else None    # noqa: E731
        ws.append([x["index"], x["name"], rnd(x["value"]), x["coefficient"], rnd(x["contribution"]),
                   rnd(x["push_vs_neutral"]), x["note"]])
    _finish(ws)

    # ---------------------------------------------------------------- checklist
    hdr = [t("rep.check.col.no", lang), t("rep.check.col.item", lang), t("rep.check.col.year", lang),
           t("rep.check.col.value", lang), t("inputs.col.pages", lang), t("rep.check.col.why", lang),
           t("xl.col.source_line", lang), t("xl.col.answer", lang)]
    ws = _sheet(wb, t("xl.sheet.checklist", lang), hdr, [6, 40, 12, 22, 12, 70, 60, 24])
    dv = DataValidation(type="list", formula1='"' + t("confirm.OK", lang) + "," + t("confirm.BAD", lang) + '"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    for x in report.checklist_rows(r, lang):
        ws.append([x["no"], x["item"], x["year"], x["value"], x["pages"], x["why"], x["quote"], ""])
        dv.add(ws.cell(row=ws.max_row, column=8))
        if x["assumed"]:
            for c in ws[ws.max_row]:
                c.fill = FILL[report.YELLOW]
    _finish(ws)
    ws.append([])
    ws.append(["", t("confirm.reviewer", lang)])

    # ---------------------------------------------------------------- cross-checks (failed and review first)
    col_name = lambda c: {"t": yt, "t1": yt1}.get(c) or t("col.both", lang, a=yt, b=yt1)   # noqa: E731
    ws = _sheet(wb, t("xl.sheet.checks", lang), [t("rep.b.col.no", lang), t("rep.b.col.check", lang),
                                                  t("rep.b.col.result", lang), t("rep.b.col.expl", lang)], [6, 60, 16, 70])
    failed, review, passed = [], [], []
    for name_, col, ok, det in ex.checks:
        chk, d = f"{render(name_, lang)} ({col_name(col)})", render(det, lang)
        (passed if ok else failed).append((chk, d if ok else t("rep.b.failed_why", lang) + (" " + d + "." if d else "")))
    for p_ in r.passed:
        passed.append((render(p_, lang), ""))
    for fl in r.flags:
        fields = ", ".join(report.title(x, lang) for x in fl.fields) or "—"
        (failed if fl.severity == "ERROR" else review).append((f"{fields} ({col_name(fl.column)})", render(fl.message, lang)))
    n = 0
    for key, group in (("rep.checks.failed", failed), ("rep.checks.review", review), ("rep.checks.passed", passed)):
        for chk, expl in group:
            n += 1
            ws.append([n, chk, t(key, lang), expl])
    _finish(ws)
    if ex.issues:
        ws = _sheet(wb, t("rep.s8", lang).split("·")[-1].strip(), [t("rep.b.col.no", lang), t("rep.c.col.note", lang)], [6, 120])
        for k, i_ in enumerate(ex.issues, 1):
            ws.append([k, render(i_, lang)])
        _finish(ws)

    for ws in wb.worksheets:              # PDF text layers can carry control characters (NRCA 2021)
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    c.value = ILLEGAL_CHARACTERS_RE.sub("", c.value).replace("\ufffe", "").replace("\uffff", "")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
