"""Provisional result (Gail, 28 Oct 2026), issuer identification, source labels and the downloads.

* The provisional score uses only workbook statuses: it equals the score the user would get by
  choosing PARTIAL_UNDISCLOSED (M-06) and, where no intangible assets were read, amortisation 0 (M-15).
* Every assumption is marked ASSUMED and appears in the checklist; nothing else changes.
* The issuer is identified only on a unique match with the master list; otherwise not identified.
* The PDF report and the Excel table build in both languages; blank figures stay blank.
"""
import io
from pathlib import Path

import openpyxl
import pytest

from wars_scanner import company, pipeline, report, report_pdf, report_xlsx
from wars_scanner.text import LANGS

PDF_DIR = Path(__file__).resolve().parents[1].parent / "s3core" / "docs"
DOCS = ["LK_SIDO_2021.pdf", "LK_MYOR_2020.pdf", "LK_MYOR_2021.pdf"]


@pytest.fixture(scope="module")
def runs():
    out = {}
    for name in DOCS:
        p = PDF_DIR / name
        if p.exists():
            out[name] = pipeline.run(str(p), "X")
    if not out:
        pytest.skip("test documents not available")
    return out


def test_provisional_score_equals_the_workbook_choices(runs):
    for name, r in runs.items():
        prov = pipeline.apply_provisional(r)
        if r.unreadable:
            assert prov is r
            continue
        am = {c: {"status": "NONE_PRESENT", "notes": ["(test)"]} for c in r.amort_pending}
        np_ = {c: {"status": "PARTIAL_UNDISCLOSED", "notes": []} for c in ("t", "t1")}
        ref = pipeline.apply_user_decisions(r, np_, am)
        assert prov.m_score is not None, name
        assert prov.m_score == pytest.approx(ref.m_score, abs=1e-12), name
        assert prov.status == "PROVISIONAL (inputs to check)"


def test_every_assumption_is_marked_and_listed(runs):
    for name, r in runs.items():
        if r.unreadable:
            continue
        prov = pipeline.apply_provisional(r)
        expected = {("amortization_intangibles", c) for c in r.amort_pending} | \
                   {("depreciation_amortization_non_ppe", c) for c in ("t", "t1")}
        assert {(f, c) for f, c, _ in prov.assumed} == expected, name
        for f, c in expected:
            v = next(v for v in prov.extraction.values if v.field == f and v.column == c)
            assert v.status == "ASSUMED" and v.note, (name, f, c)
        listed = {(x["key"].split(":")[1], x["key"].split(":")[2]) for x in report.checklist_rows(prov) if x["assumed"]}
        assert listed == expected, name
        # values that were read are not touched
        before = {(v.field, v.column): v.value for v in r.extraction.values if v.status == "FOUND"}
        after = {(v.field, v.column): v.value for v in prov.extraction.values}
        assert all(after[k] == val for k, val in before.items()), name


def test_issuer_identification_needs_a_unique_match():
    who, _ = company.identify("PT KALBE FARMA Tbk. DAN ENTITAS ANAKNYA LAPORAN POSISI KEUANGAN")
    assert who is not None and who.ticker == "KLBF"
    who, _ = company.identify("PT INDUSTRI JAMU DAN FARMASI PT INDUSTRI JAMU DAN FARMASI SIDO MUNCUL TBK SIDO MUNCUL TBK")
    assert who is not None and who.ticker == "SIDO"
    who, cands = company.identify("LAPORAN POSISI KEUANGAN KONSOLIDASIAN")
    assert who is None
    who, _ = company.identify("")
    assert who is None


def test_source_label_follows_the_language_and_drops_figures():
    q = "PENJUALAN 4.020.980 22,31 3.335.411 SALES"
    assert report.source_label(q, "id") == "PENJUALAN"
    assert report.source_label(q, "en") == "SALES"
    q = "Beban penjualan 3.769.234.915.945 4.027.986.833.485"            # Indonesian-only line: its label
    assert report.source_label(q, "en") == "Beban penjualan"
    assert report.source_label("Pajak Penghasilan Pasal 25 9.768 3.946 Income Tax Article 25", "id") == \
        "Pajak Penghasilan Pasal 25"                                       # a number that is part of the label stays
    assert report.source_label("Pihak berelasi 6,23 2.115.494.369 1.126.732.689 Related parties", "id") == "Pihak berelasi"
    assert report.source_label("", "en") == ""


def test_downloads_build_in_both_languages(runs):
    for name, r in runs.items():
        if r.unreadable:
            continue
        prov = pipeline.apply_provisional(r)
        for lang in LANGS:
            pdf = report_pdf.build(prov, lang, name, "0" * 64)
            assert pdf[:5] == b"%PDF-" and len(pdf) > 20000
            wb = openpyxl.load_workbook(io.BytesIO(report_xlsx.build(prov, lang, name, "0" * 64, "x")))
            assert len(wb.sheetnames) == (6 if r.extraction.issues else 5)     # + reading notes
            figures = wb.worksheets[1]
            for row in figures.iter_rows(min_row=2, values_only=True):
                assert row[1] is None or isinstance(row[1], (int, float))
        break                       # one document is enough for the build; the rest is covered above


def test_blank_figure_stays_blank_in_excel(runs):
    r = next((r for r in runs.values() if not r.unreadable), None)
    if r is None:
        pytest.skip("no readable document")
    prov = pipeline.apply_provisional(r)
    blank = {report.label(v.field, "en") for v in prov.extraction.values
             if v.value is None and v.column == "t" and v.field in report.FIELD_ORDER}
    wb = openpyxl.load_workbook(io.BytesIO(report_xlsx.build(prov, "en", "x.pdf", "0" * 64, "x")))
    for row in wb.worksheets[1].iter_rows(min_row=2, values_only=True):
        if row[0].strip() in {b.strip() for b in blank}:
            assert row[1] is None, row


def test_app_pages_open_without_errors():
    from streamlit.testing.v1 import AppTest
    root = Path(__file__).resolve().parents[1]
    for lang in LANGS:
        at = AppTest.from_file(str(root / "app.py"), default_timeout=60)
        at.session_state["lang"] = lang
        at.run()
        assert not at.exception
        at.button(key="guest").click().run()
        assert not at.exception
        assert at.session_state["page"] == "upload"
